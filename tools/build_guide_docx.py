"""Render strategy_guide.md into a typeset Word document.

The markdown is the source of truth. This builder adds the presentation layer:
cover page, table of contents, running header and page numbers, styled tables,
stage-box figures, callouts, shaded command blocks, and the native Word
equations captured by ``tools/extract_equations.py``.

Usage:
    python tools/build_guide_docx.py
    python tools/build_guide_docx.py --source strategy_guide.md --output out.docx

Supported markdown subset:
    ---            YAML-style front matter (title, subtitle, author, ...)
    # / ## / ###   headings
    - item         bullet list
    1. item        ordered list
    > text         callout box
    | a | b |      table, optionally preceded by a "Table: caption" line
    ```flow        stage-box figure ("caption:", "direction: right", "Name :: note")
    ```powershell  shaded command block (any other fence language works too)
    {{eq:slug}}    native Word equation, inline or on its own line
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree

try:
    from docx import Document
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn
    from docx.shared import Inches, Pt, RGBColor
except ModuleNotFoundError:  # pragma: no cover - dependency guard
    sys.exit(
        "python-docx is required.\n"
        "Install it with: .\\.venv\\Scripts\\python.exe -m pip install -r tools/requirements-docs.txt"
    )

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SOURCE = ROOT / "strategy_guide.md"
DEFAULT_OUTPUT = ROOT / "QuantLab_Investment_Strategy_and_Research_Guide_v3.docx"
EQUATIONS = ROOT / "tools" / "guide_assets" / "equations.json"

ACCENT = "1F3864"
ACCENT_SOFT = "2E4A7D"
INK = "1A1A1A"
MUTED = "5A6472"
RULE = "D3DAE6"
BOX_FILL = "EEF2F8"
CODE_FILL = "F5F7FA"
HEADER_FILL = "1F3864"
BAND_FILL = "F4F6FA"

BODY_FONT = "Cambria"
DISPLAY_FONT = "Segoe UI"
MONO_FONT = "Consolas"

CONTENT_WIDTH = Inches(6.5)

EQ_TOKEN = re.compile(r"\{\{eq:([a-z0-9_]+)\}\}")
INLINE_TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|\{\{eq:[a-z0-9_]+\}\})")


# --------------------------------------------------------------------------
# markdown parsing
# --------------------------------------------------------------------------


@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    items: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    caption: str = ""
    language: str = ""
    direction: str = "down"


def parse_front_matter(lines: list[str]) -> tuple[dict[str, str], list[str]]:
    if not lines or lines[0].strip() != "---":
        return {}, lines
    meta: dict[str, str] = {}
    for index in range(1, len(lines)):
        line = lines[index]
        if line.strip() == "---":
            return meta, lines[index + 1 :]
        if ":" in line:
            key, _, value = line.partition(":")
            meta[key.strip()] = value.strip().strip('"')
    return meta, lines


def parse_markdown(text: str) -> tuple[dict[str, str], list[Block]]:
    meta, lines = parse_front_matter(text.splitlines())
    blocks: list[Block] = []
    pending_caption = ""
    index = 0

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        if stripped.startswith("```"):
            language = stripped[3:].strip()
            index += 1
            body: list[str] = []
            while index < len(lines) and lines[index].strip() != "```":
                body.append(lines[index])
                index += 1
            index += 1
            if language == "flow":
                blocks.append(build_flow_block(body))
            else:
                blocks.append(
                    Block(kind="code", language=language, text="\n".join(body))
                )
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            blocks.append(
                Block(kind="heading", level=level, text=stripped[level:].strip())
            )
            index += 1
            continue

        if stripped.startswith("Table:"):
            pending_caption = stripped[len("Table:") :].strip()
            index += 1
            continue

        if stripped.startswith("|"):
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [
                    cell.strip() for cell in lines[index].strip().strip("|").split("|")
                ]
                if not all(set(cell) <= set("-: ") and cell for cell in cells):
                    rows.append(cells)
                index += 1
            blocks.append(Block(kind="table", rows=rows, caption=pending_caption))
            pending_caption = ""
            continue

        if stripped.startswith("> "):
            quote: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("> "):
                quote.append(lines[index].strip()[2:])
                index += 1
            blocks.append(Block(kind="callout", text=" ".join(quote)))
            continue

        if stripped.startswith("- "):
            items = []
            while index < len(lines) and lines[index].strip().startswith("- "):
                items.append(lines[index].strip()[2:])
                index += 1
            blocks.append(Block(kind="bullets", items=items))
            continue

        if re.match(r"^\d+\.\s", stripped):
            items = []
            while index < len(lines) and re.match(r"^\d+\.\s", lines[index].strip()):
                items.append(re.sub(r"^\d+\.\s", "", lines[index].strip()))
                index += 1
            blocks.append(Block(kind="numbers", items=items))
            continue

        paragraph = [stripped]
        index += 1
        while index < len(lines) and lines[index].strip() and not is_block_start(
            lines[index]
        ):
            paragraph.append(lines[index].strip())
            index += 1
        joined = " ".join(paragraph)
        if EQ_TOKEN.fullmatch(joined):
            blocks.append(Block(kind="equation", text=EQ_TOKEN.match(joined).group(1)))
        else:
            blocks.append(Block(kind="paragraph", text=joined))

    return meta, blocks


def is_block_start(line: str) -> bool:
    stripped = line.strip()
    return (
        stripped.startswith(("#", "```", "|", "> ", "- ", "Table:"))
        or bool(re.match(r"^\d+\.\s", stripped))
    )


def build_flow_block(body: list[str]) -> Block:
    block = Block(kind="flow")
    for raw in body:
        line = raw.strip()
        if not line:
            continue
        if line.lower().startswith("caption:"):
            block.caption = line.split(":", 1)[1].strip()
        elif line.lower().startswith("direction:"):
            block.direction = line.split(":", 1)[1].strip().lower()
        else:
            block.items.append(line)
    return block


# --------------------------------------------------------------------------
# low-level docx helpers
# --------------------------------------------------------------------------


def shade(cell, fill: str) -> None:
    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>')
    )


def set_borders(cell, **edges) -> None:
    """edges: top/bottom/left/right -> (size_eighths, color) or None for 'nil'."""
    parts = []
    for edge in ("top", "left", "bottom", "right"):
        if edge not in edges:
            continue
        spec = edges[edge]
        if spec is None:
            parts.append(f'<w:{edge} w:val="nil"/>')
        else:
            size, color = spec
            parts.append(
                f'<w:{edge} w:val="single" w:sz="{size}" w:space="0" w:color="{color}"/>'
            )
    if parts:
        cell._tc.get_or_add_tcPr().append(
            parse_xml(f'<w:tcBorders {nsdecls("w")}>{"".join(parts)}</w:tcBorders>')
        )


def set_margins(cell, top=60, left=110, bottom=60, right=110) -> None:
    cell._tc.get_or_add_tcPr().append(
        parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f"</w:tcMar>"
        )
    )


def bottom_rule(paragraph, color: str = RULE, size: int = 6) -> None:
    paragraph._p.get_or_add_pPr().append(
        parse_xml(
            f'<w:pBdr {nsdecls("w")}>'
            f'<w:bottom w:val="single" w:sz="{size}" w:space="4" w:color="{color}"/>'
            f"</w:pBdr>"
        )
    )


def keep_together(paragraph) -> None:
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = True


def style_run(run, *, font=BODY_FONT, size=10.5, bold=False, italic=False, color=INK):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.rPr.rFonts.set(qn("w:cs"), font)
    return run


def field_run(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    run._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>'))
    run._r.append(
        parse_xml(
            f'<w:instrText {nsdecls("w")} xml:space="preserve"> {instruction} </w:instrText>'
        )
    )
    run._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>'))
    if placeholder:
        style_run(paragraph.add_run(placeholder), size=9, color=MUTED)
    end = paragraph.add_run()
    end._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))


# --------------------------------------------------------------------------
# builder
# --------------------------------------------------------------------------


class GuideBuilder:
    def __init__(self, meta: dict[str, str], equations: dict[str, str]):
        self.meta = meta
        self.equations = equations
        self.document = Document()
        self.expected_text: list[str] = []
        self.equation_uses = 0
        self.table_number = 0
        self.figure_number = 0
        self._configure_styles()
        self._configure_page()
        self._configure_properties()

    # -- setup ------------------------------------------------------------

    def _configure_properties(self) -> None:
        properties = self.document.core_properties
        properties.title = self.meta.get("title", "")
        properties.subject = self.meta.get("subtitle", "")
        properties.author = self.meta.get("author", "")
        properties.last_modified_by = self.meta.get("author", "")
        properties.keywords = self.meta.get("keywords", "")
        properties.category = self.meta.get("status", "")
        properties.comments = self.meta.get("notice", "")

    def _configure_page(self) -> None:
        section = self.document.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        for attribute in ("top_margin", "bottom_margin"):
            setattr(section, attribute, Inches(0.9))
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.different_first_page_header_footer = True

        header = section.header.paragraphs[0]
        header.text = ""
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        style_run(
            header.add_run(self.meta.get("title", "QuantLab")),
            font=DISPLAY_FONT,
            size=8,
            color=MUTED,
        )
        bottom_rule(header)

        # A two-column table rather than tab stops: tab rendering varies between
        # Word and WPS, table cells do not.
        footer_table = section.footer.add_table(
            rows=1, cols=2, width=CONTENT_WIDTH
        )
        footer_table.autofit = False
        notice_cell, page_cell = footer_table.rows[0].cells
        notice_cell.width = Inches(4.5)
        page_cell.width = Inches(2.0)
        for cell in (notice_cell, page_cell):
            set_borders(cell, top=None, left=None, bottom=None, right=None)
            set_margins(cell, top=0, bottom=0, left=0, right=0)

        notice = self.cell_paragraph(notice_cell)
        style_run(
            notice.add_run("Research document \u2014 not an investment recommendation"),
            font=DISPLAY_FONT,
            size=8,
            color=MUTED,
        )

        pages = self.cell_paragraph(page_cell)
        pages.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        style_run(pages.add_run("Page "), font=DISPLAY_FONT, size=8, color=MUTED)
        field_run(pages, "PAGE")
        style_run(pages.add_run(" of "), font=DISPLAY_FONT, size=8, color=MUTED)
        field_run(pages, "NUMPAGES")

        # The footer must still end with a paragraph, so move the empty one below.
        original = section.footer.paragraphs[0]._p
        original.getparent().remove(original)
        closing = section.footer.add_paragraph()
        closing.paragraph_format.space_before = Pt(0)
        closing.paragraph_format.space_after = Pt(0)
        closing.paragraph_format.line_spacing = 1.0
        style_run(closing.add_run(""), size=1)

    def _configure_styles(self) -> None:
        styles = self.document.styles

        normal = styles["Normal"]
        normal.font.name = BODY_FONT
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = RGBColor.from_string(INK)
        normal.paragraph_format.space_after = Pt(8)
        normal.paragraph_format.line_spacing = 1.15

        heading_specs = {
            "Heading 1": (15, ACCENT, 22, 8),
            "Heading 2": (12, ACCENT_SOFT, 16, 6),
            "Heading 3": (10.5, ACCENT_SOFT, 12, 4),
        }
        for name, (size, color, before, after) in heading_specs.items():
            style = styles[name]
            style.font.name = DISPLAY_FONT
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.italic = False
            style.font.color.rgb = RGBColor.from_string(color)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.line_spacing = 1.0

        # Update the table of contents automatically when Word opens the file.
        self.document.settings.element.append(
            parse_xml(f'<w:updateFields {nsdecls("w")} w:val="true"/>')
        )

    # -- primitives -------------------------------------------------------

    def paragraph(self, text: str, **kwargs):
        paragraph = self.document.add_paragraph()
        self.write_inline(paragraph, text, **kwargs)
        return paragraph

    def write_inline(self, paragraph, text: str, *, size=10.5, color=INK, italic=False):
        """Render bold/italic/code/equation markup into an existing paragraph."""
        literal: list[str] = []
        for piece in INLINE_TOKEN.split(text):
            if not piece:
                continue
            equation = EQ_TOKEN.fullmatch(piece)
            if equation:
                self.append_equation(paragraph, equation.group(1))
            elif piece.startswith("**") and piece.endswith("**"):
                content = typeset(piece[2:-2])
                style_run(
                    paragraph.add_run(content),
                    size=size,
                    bold=True,
                    italic=italic,
                    color=color,
                )
                literal.append(content)
            elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
                content = typeset(piece[1:-1])
                style_run(
                    paragraph.add_run(content), size=size, italic=True, color=color
                )
                literal.append(content)
            elif piece.startswith("`") and piece.endswith("`"):
                # Code keeps straight quotes: these strings get copied into a shell.
                style_run(
                    paragraph.add_run(piece[1:-1]),
                    font=MONO_FONT,
                    size=size - 1,
                    color=ACCENT_SOFT,
                )
                literal.append(piece[1:-1])
            else:
                content = typeset(piece)
                style_run(
                    paragraph.add_run(content), size=size, italic=italic, color=color
                )
                literal.append(content)
        joined = normalize("".join(literal))
        if joined:
            self.expected_text.append(joined)
        return paragraph

    def append_equation(self, paragraph, slug: str) -> None:
        fragment = self.equations.get(slug)
        if fragment is None:
            raise KeyError(f"unknown equation slug: {slug}")
        paragraph._p.append(parse_xml(fragment))
        self.equation_uses += 1

    def caption(self, label: str, text: str, *, above: bool) -> None:
        paragraph = self.document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(2 if above else 4)
        paragraph.paragraph_format.space_after = Pt(4 if above else 10)
        paragraph.paragraph_format.keep_with_next = above
        style_run(
            paragraph.add_run(f"{label} "),
            font=DISPLAY_FONT,
            size=8.5,
            bold=True,
            color=ACCENT,
        )
        self.write_inline(paragraph, text, size=8.5, color=MUTED)
        for run in paragraph.runs[1:]:
            if run.font.name != MONO_FONT:
                run.font.name = DISPLAY_FONT

    def new_table(self, rows: int, cols: int):
        table = self.document.add_table(rows=rows, cols=cols)
        table.autofit = False
        table.alignment = WD_ALIGN_PARAGRAPH.CENTER
        table._tbl.tblPr.append(
            parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>')
        )
        return table

    @staticmethod
    def cell_paragraph(cell, *, first: bool = True):
        paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        return paragraph

    # -- block renderers --------------------------------------------------

    def render_heading(self, block: Block) -> None:
        heading = self.document.add_heading(level=min(block.level, 3))
        for run in list(heading.runs):
            run._element.getparent().remove(run._element)
        style_run(
            heading.add_run(block.text),
            font=DISPLAY_FONT,
            size={1: 15, 2: 12, 3: 10.5}[min(block.level, 3)],
            bold=True,
            color=ACCENT if block.level == 1 else ACCENT_SOFT,
        )
        if block.level == 1:
            bottom_rule(heading, color=RULE, size=8)
        self.expected_text.append(normalize(block.text))

    def render_list(self, block: Block, ordered: bool) -> None:
        for position, item in enumerate(block.items, start=1):
            paragraph = self.document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.32)
            paragraph.paragraph_format.first_line_indent = Inches(-0.22)
            paragraph.paragraph_format.space_after = Pt(3)
            marker = f"{position}." if ordered else "\u2022"
            style_run(
                paragraph.add_run(f"{marker}\t"),
                font=DISPLAY_FONT if ordered else BODY_FONT,
                size=10.5,
                bold=ordered,
                color=ACCENT_SOFT if ordered else ACCENT,
            )
            self.write_inline(paragraph, item)

    def render_table(self, block: Block) -> None:
        if not block.rows:
            return
        self.table_number += 1
        if block.caption:
            self.caption(f"Table {self.table_number}.", block.caption, above=True)

        columns = len(block.rows[0])
        table = self.new_table(rows=len(block.rows), cols=columns)
        widths = column_widths(block.rows)
        for index, width in enumerate(widths):
            table.columns[index].width = width

        for row_index, row in enumerate(block.rows):
            is_header = row_index == 0
            row_properties = table.rows[row_index]._tr.get_or_add_trPr()
            row_properties.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            if is_header:
                # Repeat the header when a table runs over a page break.
                row_properties.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            for column_index in range(columns):
                cell = table.cell(row_index, column_index)
                cell.width = widths[column_index]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_margins(cell, top=70, bottom=70)
                set_borders(
                    cell,
                    top=None,
                    left=None,
                    right=None,
                    bottom=(8, ACCENT) if is_header else (4, RULE),
                )
                if is_header:
                    shade(cell, HEADER_FILL)
                elif row_index % 2 == 0:
                    shade(cell, BAND_FILL)

                text = row[column_index] if column_index < len(row) else ""
                paragraph = self.cell_paragraph(cell)
                self.write_inline(
                    paragraph,
                    text,
                    size=9,
                    color="FFFFFF" if is_header else INK,
                )
                for run in paragraph.runs:
                    run.font.name = DISPLAY_FONT
                    if is_header:
                        run.bold = True
        self.document.add_paragraph().paragraph_format.space_after = Pt(2)

    def render_flow(self, block: Block) -> None:
        stages = [
            (item.split("::")[0].strip(), item.split("::", 1)[1].strip() if "::" in item else "")
            for item in block.items
        ]
        if block.direction.startswith("r"):
            self.render_flow_horizontal(stages)
        else:
            self.render_flow_vertical(stages)
        self.figure_number += 1
        if block.caption:
            self.caption(f"Figure {self.figure_number}.", block.caption, above=False)

    def render_flow_vertical(self, stages: list[tuple[str, str]]) -> None:
        table = self.new_table(rows=len(stages) * 2 - 1, cols=1)
        for index, (name, note) in enumerate(stages):
            row_index = index * 2
            cell = table.cell(row_index, 0)
            cell.width = CONTENT_WIDTH
            shade(cell, BOX_FILL)
            set_margins(cell, top=90, bottom=90, left=170, right=170)
            set_borders(
                cell,
                left=(24, ACCENT),
                top=(4, "E3E9F2"),
                bottom=(4, "E3E9F2"),
                right=(4, "E3E9F2"),
            )
            paragraph = self.cell_paragraph(cell)
            style_run(
                paragraph.add_run(name),
                font=DISPLAY_FONT,
                size=10,
                bold=True,
                color=ACCENT,
            )
            self.expected_text.append(normalize(name))
            if note:
                note_paragraph = self.cell_paragraph(cell, first=False)
                style_run(
                    note_paragraph.add_run(note),
                    font=DISPLAY_FONT,
                    size=8.5,
                    color=MUTED,
                )
                self.expected_text.append(normalize(note))

            if index < len(stages) - 1:
                arrow_cell = table.cell(row_index + 1, 0)
                arrow_cell.width = CONTENT_WIDTH
                set_margins(arrow_cell, top=20, bottom=20)
                set_borders(arrow_cell, top=None, left=None, bottom=None, right=None)
                arrow = self.cell_paragraph(arrow_cell)
                arrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
                style_run(
                    arrow.add_run("\u25bc"), font=DISPLAY_FONT, size=9, color=ACCENT_SOFT
                )

    def render_flow_horizontal(self, stages: list[tuple[str, str]]) -> None:
        columns = len(stages) * 2 - 1
        table = self.new_table(rows=1, cols=columns)
        arrow_width = Inches(0.28)
        stage_width = int(
            (CONTENT_WIDTH.emu - arrow_width.emu * (len(stages) - 1)) / len(stages)
        )

        for index, (name, note) in enumerate(stages):
            cell = table.cell(0, index * 2)
            cell.width = stage_width
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            shade(cell, BOX_FILL)
            set_margins(cell, top=90, bottom=90, left=90, right=90)
            set_borders(
                cell,
                top=(4, "E3E9F2"),
                left=(4, "E3E9F2"),
                bottom=(12, ACCENT),
                right=(4, "E3E9F2"),
            )
            paragraph = self.cell_paragraph(cell)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style_run(
                paragraph.add_run(name),
                font=DISPLAY_FONT,
                size=9,
                bold=True,
                color=ACCENT,
            )
            self.expected_text.append(normalize(name))
            if note:
                note_paragraph = self.cell_paragraph(cell, first=False)
                note_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                style_run(
                    note_paragraph.add_run(note),
                    font=DISPLAY_FONT,
                    size=8,
                    color=MUTED,
                )
                self.expected_text.append(normalize(note))

            if index < len(stages) - 1:
                arrow_cell = table.cell(0, index * 2 + 1)
                arrow_cell.width = arrow_width
                arrow_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_borders(arrow_cell, top=None, left=None, bottom=None, right=None)
                arrow = self.cell_paragraph(arrow_cell)
                arrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
                style_run(
                    arrow.add_run("\u25b6"),
                    font=DISPLAY_FONT,
                    size=9,
                    color=ACCENT_SOFT,
                )

    def render_code(self, block: Block) -> None:
        table = self.new_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        cell.width = CONTENT_WIDTH
        shade(cell, CODE_FILL)
        set_margins(cell, top=110, bottom=110, left=150, right=150)
        set_borders(
            cell,
            top=(4, RULE),
            left=(18, ACCENT_SOFT),
            bottom=(4, RULE),
            right=(4, RULE),
        )
        for index, line in enumerate(block.text.splitlines()):
            paragraph = self.cell_paragraph(cell, first=index == 0)
            keep_together(paragraph)
            style_run(
                paragraph.add_run(line or " "),
                font=MONO_FONT,
                size=8.5,
                color="102A54",
            )
            if line.strip():
                self.expected_text.append(normalize(line))
        self.document.add_paragraph().paragraph_format.space_after = Pt(2)

    def render_callout(self, block: Block) -> None:
        table = self.new_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        cell.width = CONTENT_WIDTH
        shade(cell, BOX_FILL)
        set_margins(cell, top=130, bottom=130, left=170, right=170)
        set_borders(
            cell, top=None, left=(30, ACCENT), bottom=None, right=None
        )
        paragraph = self.cell_paragraph(cell)
        self.write_inline(paragraph, block.text, size=11, color=ACCENT, italic=True)
        self.document.add_paragraph().paragraph_format.space_after = Pt(2)

    def render_equation(self, slug: str) -> None:
        paragraph = self.document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(10)
        self.append_equation(paragraph, slug)

    # -- document assembly ------------------------------------------------

    def build_cover(self) -> None:
        spacer = self.document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(150)

        eyebrow = self.document.add_paragraph()
        eyebrow.paragraph_format.space_after = Pt(10)
        style_run(
            eyebrow.add_run(self.meta.get("status", "").upper()),
            font=DISPLAY_FONT,
            size=9,
            bold=True,
            color=ACCENT_SOFT,
        )

        title = self.document.add_paragraph()
        title.paragraph_format.space_after = Pt(6)
        style_run(
            title.add_run(self.meta.get("title", "")),
            font=DISPLAY_FONT,
            size=26,
            bold=True,
            color=ACCENT,
        )

        subtitle = self.document.add_paragraph()
        subtitle.paragraph_format.space_after = Pt(14)
        style_run(
            subtitle.add_run(self.meta.get("subtitle", "")),
            font=DISPLAY_FONT,
            size=13,
            color=MUTED,
        )
        bottom_rule(subtitle, color=ACCENT, size=12)

        details = self.document.add_paragraph()
        details.paragraph_format.space_after = Pt(4)
        style_run(
            details.add_run(self.meta.get("author", "")),
            font=DISPLAY_FONT,
            size=11,
            bold=True,
            color=INK,
        )
        details.add_run("\n")
        style_run(
            details.add_run(self.meta.get("date", "")),
            font=DISPLAY_FONT,
            size=10,
            color=MUTED,
        )

        keywords = self.document.add_paragraph()
        keywords.paragraph_format.space_before = Pt(220)
        style_run(
            keywords.add_run(self.meta.get("keywords", "")),
            font=DISPLAY_FONT,
            size=8.5,
            color=MUTED,
        )

        if self.meta.get("notice"):
            self.render_callout(Block(kind="callout", text=self.meta["notice"]))

        self.document.add_page_break()

    def build_toc(self) -> None:
        heading = self.document.add_paragraph()
        heading.paragraph_format.space_after = Pt(10)
        style_run(
            heading.add_run("Contents"),
            font=DISPLAY_FONT,
            size=15,
            bold=True,
            color=ACCENT,
        )
        bottom_rule(heading, color=RULE, size=8)

        paragraph = self.document.add_paragraph()
        field_run(
            paragraph,
            'TOC \\o "1-3" \\h \\z \\u',
            "Right-click here and choose Update Field to build the contents.",
        )
        self.document.add_page_break()

    def build_body(self, blocks: list[Block]) -> None:
        renderers = {
            "heading": self.render_heading,
            "paragraph": lambda block: self.paragraph(block.text),
            "bullets": lambda block: self.render_list(block, ordered=False),
            "numbers": lambda block: self.render_list(block, ordered=True),
            "table": self.render_table,
            "flow": self.render_flow,
            "code": self.render_code,
            "callout": self.render_callout,
            "equation": lambda block: self.render_equation(block.text),
        }
        for block in blocks:
            renderers[block.kind](block)

    def save(self, path: Path) -> None:
        self.document.save(path)


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def column_widths(rows: list[list[str]], minimum: float = 0.8) -> list[int]:
    """Share the text column between table columns in proportion to their content.

    Equal columns leave identifiers such as ``factors.StockExposureSnapshot``
    breaking mid-word while a one-word column keeps space it cannot use.
    """
    columns = len(rows[0])
    plain = [
        [normalize(re.sub(r"[*`]|\{\{eq:[a-z0-9_]+\}\}", "", cell)) for cell in row]
        for row in rows
    ]
    demand = []
    for index in range(columns):
        cells = [row[index] for row in plain if index < len(row)]
        longest_line = min(max((len(cell) for cell in cells), default=1), 80)
        longest_word = max(
            (len(word) for cell in cells for word in cell.split()), default=1
        )
        demand.append(max(0.45 * longest_line, 0.9 * longest_word, 4.0))

    total_inches = CONTENT_WIDTH.inches
    scale = total_inches / sum(demand)
    widths = [value * scale for value in demand]

    # Apply the floor, then take the difference back from the columns that can spare it.
    deficit = sum(minimum - width for width in widths if width < minimum)
    if deficit:
        spare = [width for width in widths if width > minimum]
        reducible = sum(spare) - minimum * len(spare)
        widths = [
            minimum
            if width < minimum
            else width - deficit * (width - minimum) / reducible
            for width in widths
        ]
    return [Inches(width) for width in widths]


def typeset(text: str) -> str:
    """Straight quotes are convenient in markdown; Word should show real ones."""
    text = re.sub(r"(?<=\w)'(?=\w)", "\u2019", text)
    text = re.sub(r"(?<=\w)'(?=\W|$)", "\u2019", text)
    result: list[str] = []
    opening = True
    for character in text:
        if character == '"':
            result.append("\u201c" if opening else "\u201d")
            opening = not opening
        else:
            result.append(character)
    return "".join(result)


# --------------------------------------------------------------------------
# verification
# --------------------------------------------------------------------------


def verify(path: Path, builder: GuideBuilder, expected_equations: int) -> list[str]:
    """Read the saved file back and confirm nothing was dropped in rendering."""
    with zipfile.ZipFile(path) as archive:
        document = archive.read("word/document.xml")

    problems: list[str] = []

    found = document.count(b"<m:oMath")
    if found != expected_equations:
        problems.append(
            f"expected {expected_equations} equations in output, found {found}"
        )

    root = ElementTree.fromstring(document)
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    # One line per paragraph, so a match can never span two paragraphs.
    rendered = "\n".join(
        normalize("".join(node.text or "" for node in paragraph.iter(w + "t")))
        for paragraph in root.iter(w + "p")
    )

    for text in builder.expected_text:
        if text not in rendered:
            problems.append(f"text missing from output: {text[:90]!r}")

    return problems


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    meta, blocks = parse_markdown(args.source.read_text(encoding="utf-8"))
    equations = json.loads(EQUATIONS.read_text(encoding="utf-8"))

    builder = GuideBuilder(meta, equations)
    builder.build_cover()
    builder.build_toc()
    builder.build_body(blocks)
    builder.save(args.output)

    expected_equations = len(EQ_TOKEN.findall(args.source.read_text(encoding="utf-8")))
    problems = verify(args.output, builder, expected_equations)

    print(f"source     {args.source.name}")
    print(f"output     {args.output.name}")
    print(
        f"rendered   {builder.table_number} tables, {builder.figure_number} figures, "
        f"{builder.equation_uses} equations"
    )
    if problems:
        print(f"FAILED     {len(problems)} verification problem(s):")
        for problem in problems[:20]:
            print(f"  - {problem}")
        return 1
    print("verified   every markdown text unit and equation is present in the docx")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
